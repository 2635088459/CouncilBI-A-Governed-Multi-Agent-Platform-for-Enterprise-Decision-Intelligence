"""Structured file parsing: CSV/TSV/XLSX to inferred schema plus row data.

Column names and types are inferred by DuckDB itself (FR-FV10-008), not by a
hand-rolled type guesser: XLSX is first flattened to CSV bytes from its first
sheet, then both formats go through the same ``read_csv_auto`` scan so a
single inference path backs every structured format.
"""

from __future__ import annotations

import csv
import io
import tempfile
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Mapping

import duckdb
import openpyxl


MAX_STRUCTURED_FILE_ROWS = 1_000_000


class RowLimitExceeded(Exception):
    def __init__(self, row_count: int) -> None:
        self.row_count = row_count
        super().__init__(
            f"row_count {row_count} exceeds the {MAX_STRUCTURED_FILE_ROWS} row limit"
        )


@dataclass(frozen=True, slots=True)
class ColumnSchema:
    name: str
    type: str


@dataclass(frozen=True, slots=True)
class ParsedTable:
    columns: tuple[ColumnSchema, ...]
    rows: tuple[Mapping[str, Any], ...]
    row_count: int


class StructuredFileParser:
    """Parse structured uploads into an inferred schema and row set via DuckDB."""

    def parse(self, content_bytes: bytes, extension: str) -> ParsedTable:
        normalized_extension = extension.lower().lstrip(".")
        if normalized_extension == "csv":
            return self.parse_csv(content_bytes, delimiter=",")
        if normalized_extension == "tsv":
            return self.parse_csv(content_bytes, delimiter="\t")
        if normalized_extension == "xlsx":
            return self.parse_xlsx(content_bytes)
        raise ValueError(f"unsupported structured file extension: {extension}")

    def parse_csv(self, content_bytes: bytes, *, delimiter: str = ",") -> ParsedTable:
        _enforce_row_limit(_count_csv_rows(content_bytes, delimiter=delimiter))
        return _infer_via_duckdb(content_bytes, delimiter=delimiter)

    def parse_xlsx(self, content_bytes: bytes) -> ParsedTable:
        csv_bytes = _first_sheet_to_csv_bytes(content_bytes)
        _enforce_row_limit(_count_csv_rows(csv_bytes, delimiter=","))
        return _infer_via_duckdb(csv_bytes, delimiter=",")


class ParquetWriter:
    """Write a ``ParsedTable`` to a Parquet snapshot DuckDB can scan back."""

    def write(self, table: ParsedTable, output_path: str | Path) -> None:
        connection = duckdb.connect(":memory:")
        try:
            column_defs = ", ".join(f'"{column.name}" {column.type}' for column in table.columns)
            connection.execute(f"CREATE TABLE snapshot ({column_defs})")
            if table.rows:
                placeholders = ", ".join("?" for _ in table.columns)
                row_values = [
                    tuple(row.get(column.name) for column in table.columns)
                    for row in table.rows
                ]
                connection.executemany(
                    f"INSERT INTO snapshot VALUES ({placeholders})", row_values
                )
            connection.execute(
                "COPY snapshot TO ? (FORMAT PARQUET)", [str(output_path)]
            )
        finally:
            connection.close()


# 10-followups/11: value samples let FileDataAgent's SQL-generation prompt
# show a VARCHAR column's actual stored format (e.g. '2026-06' vs 'June'),
# not just its name and type. Computed once here, from the fully-parsed
# table already in memory during upload processing — never re-derived
# per chat query (NFR-FV10-028).
SAMPLE_CARDINALITY_THRESHOLD = 20
SAMPLE_SIZE = 5


class SchemaSerializer:
    """Serialize a ``ParsedTable`` schema into the ``schema_json`` shape."""

    def to_json(self, table: ParsedTable) -> dict[str, Any]:
        return {
            "columns": [self._column_json(column, table.rows) for column in table.columns]
        }

    def _column_json(
        self, column: ColumnSchema, rows: tuple[Mapping[str, Any], ...]
    ) -> dict[str, Any]:
        entry: dict[str, Any] = {"name": column.name, "type": column.type}
        if column.type != "VARCHAR":
            return entry
        distinct = sorted({row[column.name] for row in rows if row.get(column.name) is not None})
        if not distinct:
            return entry
        if len(distinct) <= SAMPLE_CARDINALITY_THRESHOLD:
            entry["sample_values"] = distinct[:SAMPLE_SIZE]
        else:
            entry["sample_range"] = [distinct[0], distinct[-1]]
        return entry


def _enforce_row_limit(row_count: int) -> None:
    if row_count > MAX_STRUCTURED_FILE_ROWS:
        raise RowLimitExceeded(row_count)


def _count_csv_rows(content_bytes: bytes, *, delimiter: str) -> int:
    text_stream = io.StringIO(content_bytes.decode("utf-8"))
    reader = csv.reader(text_stream, delimiter=delimiter)
    data_row_count = sum(1 for _ in reader) - 1
    return max(data_row_count, 0)


def _first_sheet_to_csv_bytes(content_bytes: bytes) -> bytes:
    workbook = openpyxl.load_workbook(io.BytesIO(content_bytes), read_only=True, data_only=True)
    try:
        worksheet = workbook.worksheets[0]
        buffer = io.StringIO()
        writer = csv.writer(buffer)
        for row in worksheet.iter_rows(values_only=True):
            writer.writerow(["" if value is None else value for value in row])
        return buffer.getvalue().encode("utf-8")
    finally:
        workbook.close()


def _infer_via_duckdb(content_bytes: bytes, *, delimiter: str) -> ParsedTable:
    with tempfile.NamedTemporaryFile(suffix=".csv", delete=False) as temp_file:
        temp_file.write(content_bytes)
        temp_path = temp_file.name

    connection = duckdb.connect(":memory:")
    try:
        relation = connection.sql(
            "SELECT * FROM read_csv_auto(?, delim=?, header=true, sample_size=-1)",
            params=[temp_path, delimiter],
        )
        column_names = relation.columns
        column_types = [str(column_type) for column_type in relation.types]
        fetched_rows = relation.fetchall()
    finally:
        connection.close()
        Path(temp_path).unlink(missing_ok=True)

    columns = tuple(
        ColumnSchema(name=name, type=type_name)
        for name, type_name in zip(column_names, column_types)
    )
    rows = tuple(dict(zip(column_names, row)) for row in fetched_rows)
    return ParsedTable(columns=columns, rows=rows, row_count=len(rows))
